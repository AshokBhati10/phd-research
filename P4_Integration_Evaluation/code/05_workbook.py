"""P4 STEP 5 -- Concise integration workbook (P1/P2/P3 conventions)."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common_p4 as C

OUT = C.BASE / "results"
WB_PATH = C.BASE / "workbook" / "Person_4_Integration_Workbook.xlsx"

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F5496")

FAIRNESS = [
    ("Controlled (identical across P1/P2/P3)", "Dataset: corrected BRFSS 2015-2024, 2,297,168 rows; YEAR 2014 excluded."),
    ("Controlled", "Train/test split: stratified 80/20, random_state=42 -> 1,837,734 / 459,434; identical class counts."),
    ("Controlled", "Test set integrity: SHA-256 fingerprint 70e180c6... verified independently by P3 and P2, and re-verified by P4 from P1's split files."),
    ("Controlled", "15 int64 integer-coded predictors (no one-hot, no scaling before resampling); YEAR is metadata only."),
    ("Controlled", "Classifiers: LogisticRegression (lbfgs, max_iter=1000), RandomForest (100 trees, depth 12), XGBoost (300/6/0.05) - identical hyperparameters."),
    ("Controlled", "Metrics: Accuracy, Precision, Recall, F1, ROC-AUC, PR-AUC, CM counts; threshold 0.5; seed 42 everywhere."),
    ("Deviation 1", "P2's SVM-SMOTE and SMOTE-N use a stratified 50k train subsample (O(n^2) cost); marked with a dagger. Not strictly like-for-like vs full-train techniques."),
    ("Deviation 2", "KMeans-SMOTE needs non-default parameters on this all-categorical feature space (defaults raise RuntimeError)."),
    ("Deviation 3", "SMOTE-N generates only observed category codes (majority vote), not interpolation - inherent to the method."),
    ("Deviation 4", "SVM-SMOTE extrapolates (out_step=0.5): 2.49% of its synthetic values fall one code below the observed minimum."),
    ("Deviation 5", "XGBoost hist training is not bit-reproducible across machines/thread counts. P1 vs P2 baseline replication: LR/RF bit-identical, XGB ~1.9e-3 metric drift. Treat XGBoost gaps below ~2e-3 as cross-machine noise."),
    ("Deviation 6", "P3's SMOTE-ENN used a staged two-stage implementation (SMOTE then ENN with checkpointing) - documented as scientifically equivalent."),
    ("Deviation 7", "Probability storage differs: float32 (P1, P2) vs float64 (P3). Metrics recomputed from both match their CSVs; float32 limits recomputation agreement to ~3e-8."),
    ("Deviation 8", "Timings are machine-dependent (P1/P3 on macOS, P2 on Windows) and are excluded from comparisons."),
]

OVERVIEW = [
    ("What this workbook is", "Integration of Person 1 (basic oversampling), Person 2 (advanced oversampling) and Person 3 (cleaning & hybrid) on one identical frozen test set (n=459,434)."),
    ("How many numbers", "36 result rows = 12 techniques (1 baseline + 11 resampling strategies) x 3 classifiers. Every row is copied from a validated source artifact (see Traceability)."),
    ("Where rows come from", "P1: combined_results_2015_2024.csv (its corrected 2015-2024 audit). P2: results_person2.csv. P3: results_person3.csv. Baseline = P1's No Balancing; P2's baseline run is an independent replication (LR/RF bit-identical) and is documented, not double-counted."),
    ("Fair comparison basis", "Same dataset, same split (fingerprint-verified), same features, same classifiers, same metrics, seed 42; resampling applied to training data only."),
    ("Known deviations", "See Fairness_register sheet. The main caveat: SVM-SMOTE and SMOTE-N (dagger) trained on a 50k subsample."),
    ("Statistics", "Statistical_tests: paired McNemar on identical test rows, Holm-corrected within each classifier (66 pairs x 3 classifiers = 198 tests). AUC_bootstrap: paired bootstrap 95% CIs for champion-vs-rival AUC gaps (500 resamples, seed 42)."),
    ("How to read the metrics", "Accuracy: overall correctness (dominated by the 83% majority). Recall: share of actual diabetes cases detected - the clinically critical metric. F1: balance of precision and recall on the minority class. ROC-AUC / PR-AUC: ranking quality across all thresholds."),
    ("Headline finding", "Resampling roughly doubles minority F1 (0.15-0.22 -> 0.39-0.45). Plain Random Oversampling achieves the best F1 for all three classifiers; advanced variants add no significant F1 advantage but shape the precision/recall operating point differently, and cleaning-only methods preserve baseline ranking."),
]


def sheet(wb, name, df, widths=None):
    ws = wb.create_sheet(title=name)
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j, value=str(col))
        cell.font = HDR
        cell.fill = FILL
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, float):
                v = round(v, 10)
            ws.cell(row=i, column=j, value=v)
    for j, col in enumerate(df.columns, 1):
        if widths and col in widths:
            w = widths[col]
        else:
            body = df[col].astype(str).head(60).map(len).max() if len(df) else 0
            w = min(46, max(len(str(col)), body) + 2)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def kv_sheet(wb, name, rows, w1=26, w2=120):
    df = pd.DataFrame(rows, columns=["Item", "Description"])
    sheet(wb, name, df, widths={"Item": w1, "Description": w2})


def read(name):
    return pd.read_csv(OUT / name)


def main():
    wb = Workbook()
    wb.remove(wb.active)

    kv_sheet(wb, "Overview", OVERVIEW)

    comb = read("combined_results.csv")
    main_cols = ["Family", "Person", "Technique", "Model", "Accuracy", "Precision", "Recall",
                 "F1", "ROC_AUC", "PR_AUC", "TN", "FP", "FN", "TP", "Train_source", "Caveat",
                 "Source_results", "Source_predictions"]
    sheet(wb, "Combined_results", comb[main_cols],
          widths={"Technique": 20, "Train_source": 34, "Caveat": 30,
                  "Source_results": 30, "Source_predictions": 30})

    sheet(wb, "Family_summary", read("family_summary.csv"), widths={"Family": 24})
    sheet(wb, "Winners_tradeoffs", read("winners_tradeoffs.csv"), widths={"Winner": 24, "Runner_up": 24})
    sheet(wb, "Mean_ranks", read("mean_ranks.csv"))

    mdf = read("statistical_tests_mcnemar.csv")
    sdf = read("statistical_tests_summary.csv")
    sheet(wb, "Stat_tests_summary", sdf, widths={"F1_champion": 22, "Family_champions": 60,
                                                "Baseline_recalls_beaten": 60})
    sheet(wb, "Stat_tests_full", mdf)

    boot = OUT / "auc_bootstrap_cis.csv"
    if boot.exists():
        sheet(wb, "AUC_bootstrap_CIs", read("auc_bootstrap_cis.csv"))
    else:
        sheet(wb, "AUC_bootstrap_CIs", pd.DataFrame({"note": ["pending -- run 03_statistical_tests.py --bootstrap"]}))

    sheet(wb, "Resampling_effects", read("resampling_effects.csv"),
          widths={"Technique": 20, "Train_source": 40, "Class_counts_source": 26})

    kv_sheet(wb, "Fairness_register", FAIRNESS, w1=22, w2=110)

    trace = comb[["Technique", "Model", "Person", "F1", "Source_results", "Source_predictions"]].copy()
    trace["Verified_by"] = "P4 01_verify_inputs.py (metric recomputation from parquet + CM counts)"
    sheet(wb, "Traceability", trace, widths={"Source_results": 34, "Source_predictions": 34, "Verified_by": 40})

    WB_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WB_PATH)
    print(f"workbook saved: {WB_PATH} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
