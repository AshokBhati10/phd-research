"""STEP 4 -- Build the Person 2 Excel workbook (Person 1/3 sheet conventions)."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C

OUT = C.BASE / "results"
WORKBOOK = C.BASE / "workbook" / "Person_2_results_workbook.xlsx"

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F5496")
TECH_SHEET = {
    "Borderline-SMOTE": "Borderline_SMOTE",
    "KMeans-SMOTE": "KMeans_SMOTE",
    "SVM-SMOTE": "SVM_SMOTE",
    "SMOTE-N": "SMOTE_N",
}


def read_csv(name, required=True):
    p = OUT / name
    if not p.exists():
        if required:
            raise SystemExit(f"missing {p}")
        return None
    return pd.read_csv(p)


def sheet(wb, name, df, widths=None):
    ws = wb.create_sheet(title=name)
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j, value=str(col))
        cell.font = HDR
        cell.fill = FILL
        cell.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, float):
                v = round(v, 10)
            ws.cell(row=i, column=j, value=v)
    for j, col in enumerate(df.columns, 1):
        if widths and col in widths:
            w = widths[col]
        else:
            body = df[col].astype(str).head(50).map(len).max() if len(df) else 0
            w = min(42, max(len(str(col)), body) + 2)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def main():
    manifest = C.load_json(C.BASE / "metadata" / "manifest.json")
    env = C.load_json(C.BASE / "metadata" / "environment.json")
    dists = C.load_json(OUT / "class_distributions.json", {}) or {}
    audits = C.load_json(OUT / "audit.json", {}) or {}

    results = read_csv("results_person2.csv")
    baseline = read_csv("results_baseline_no_balancing.csv", required=False)
    combined = pd.concat([r for r in [baseline, results] if r is not None], ignore_index=True)
    combined.to_csv(OUT / "P2_combined_results.csv", index=False)

    wb = Workbook()
    wb.remove(wb.active)

    sheet(wb, "Combined_results", combined,
          widths={"Technique": 22, "Model": 20})
    for tech, sheet_name in TECH_SHEET.items():
        sub = results[results["Technique"] == tech]
        if len(sub):
            sheet(wb, sheet_name, sub, widths={"Technique": 22, "Model": 20})
    if baseline is not None:
        sheet(wb, "Baseline_NoBalancing", baseline, widths={"Technique": 22, "Model": 20})

    dist_rows = []
    for tech in [C.BASELINE] + C.TECHNIQUES:
        d = dists.get(tech)
        if not d:
            continue
        dist_rows.append({
            "Technique": tech,
            "Train_No_before": d["train_before"].get("0"),
            "Train_Yes_before": d["train_before"].get("1"),
            "Train_No_after": d["train_after"].get("0"),
            "Train_Yes_after": d["train_after"].get("1"),
            "Added_samples": d["added_samples"],
            "Final_train_size": d["final_train_size"],
            "Train_source": ("full corrected train split" if tech in
                             ("KMeans-SMOTE", "Borderline-SMOTE") else
                             ("raw corrected train split (no resampling)" if tech == C.BASELINE else
                              "shared stratified 50k subsample")),
        })
    sheet(wb, "Class_distributions", pd.DataFrame(dist_rows), widths={"Train_source": 38})

    proto_rows = [
        {"key": "scope", "value": manifest["scope"]},
        {"key": "source_splits", "value": manifest["source"]},
        {"key": "test_fingerprint_sha256", "value": manifest["test_fingerprint_sha256"]},
        {"key": "fingerprint_matches_person3", "value": manifest["fingerprint_match_person3"]},
        {"key": "n_train", "value": manifest["n_train"]},
        {"key": "n_test", "value": manifest["n_test"]},
        {"key": "train_counts", "value": str(manifest["train_counts"])},
        {"key": "test_counts", "value": str(manifest["test_counts"])},
        {"key": "split", "value": manifest["split_protocol"]},
        {"key": "target", "value": f"{C.TARGET} (No->0, Yes->1)"},
        {"key": "features", "value": ", ".join(C.FEATURES)},
        {"key": "feature_encoding", "value": "int64 integer-coded categoricals; no one-hot, no scaling before resampling"},
        {"key": "resampling_scope", "value": "TRAIN ONLY; test set frozen and never resampled"},
        {"key": "threshold", "value": "0.5"},
        {"key": "metrics", "value": "Accuracy, Precision, Recall, F1, ROC-AUC, PR-AUC, TN/FP/FN/TP, fit_predict_sec"},
    ]
    for tech in C.TECHNIQUES:
        proto_rows.append({"key": f"{tech}_params", "value": C.describe_resampler(tech)})
        proto_rows.append({"key": f"{tech}_train_source", "value": (
            "full corrected train split (1,837,734 rows)" if tech in ("Borderline-SMOTE", "KMeans-SMOTE")
            else "shared stratified 50k train subsample (seed 42)")})
        if tech in audits:
            val = audits[tech].get("synthetic_checks") or {}
            proto_rows.append({"key": f"{tech}_validity", "value": (
                f"out_of_range={val.get('out_of_range_values')}, "
                f"unobserved_codes={val.get('unobserved_code_values')}, "
                f"fractional={val.get('fractional_values')}")})
    proto_rows += [
        {"key": "LogisticRegression", "value": "max_iter=1000, solver=lbfgs, random_state=42; StandardScaler fitted on resampled train"},
        {"key": "RandomForest", "value": "n_estimators=100, max_depth=12, n_jobs=-1, random_state=42"},
        {"key": "XGBoost", "value": "n_estimators=300, max_depth=6, lr=0.05, subsample=0.8, colsample_bytree=0.8, tree_method=hist, eval_metric=logloss, n_jobs=-1, random_state=42"},
        {"key": "significance", "value": "paired McNemar on the frozen test set; Bonferroni-corrected over all tests"},
        {"key": "environment", "value": "; ".join(f"{k}={v}" for k, v in env.items() if k != "recorded_at_utc")},
    ]
    sheet(wb, "Protocol_config", pd.DataFrame(proto_rows), widths={"key": 34, "value": 110})

    mcn = read_csv("mcnemar_results.csv", required=False)
    if mcn is not None:
        sheet(wb, "Significance_McNemar", mcn)

    notes = [
        "Person 2 scope: advanced oversampling (Borderline-SMOTE, KMeans-SMOTE, SVM-SMOTE, SMOTE-N) on the corrected BRFSS 2015-2024 scope.",
        "Data + split identical to Person 1 final and Person 3: 1,837,734 train / 459,434 test; test SHA-256 matches Person 3's manifest.",
        "Borderline-SMOTE and KMeans-SMOTE trained on the full corrected train split; resampling is train-only.",
        "SVM-SMOTE and SMOTE-N are O(n^2) in samples/memory and use one shared stratified 50,000-row train subsample (seed 42). All models still evaluated on the identical frozen 459,434-row test set. Flagged with a dagger in P4 comparisons.",
        "KMeans-SMOTE library defaults (8 clusters, threshold 0.5) raise RuntimeError on this all-categorical feature space; kmeans_estimator=100 and cluster_balance_threshold=0.2 are the documented, approved P2 configuration.",
        "SMOTEN generates only observed category codes (majority vote per feature), which suits the all-categorical feature set.",
        "Synthetic-row validity checks: finite values, observed per-feature range, unobserved category codes, fractional values (all reported in Protocol_config / audit.json).",
        "No-Balancing rows are a self-contained Person 2 baseline that also cross-validates Person 1's corrected baseline.",
        "McNemar: paired tests per classifier on identical test rows; read p-values together with effect sizes (large n makes tiny gaps significant).",
        "Prediction parquets expose y_true + per-model pred/proba so all metrics are recomputable.",
    ]
    sheet(wb, "Notes", pd.DataFrame({"#": range(1, len(notes) + 1), "note": notes}),
          widths={"note": 120})

    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)
    C.log(f"workbook saved: {WORKBOOK} ({len(wb.sheetnames)} sheets: {wb.sheetnames})")


if __name__ == "__main__":
    main()
